from __future__ import annotations

import bisect
import csv
import colorsys
import math
import tkinter as tk
from concurrent.futures import Future, ThreadPoolExecutor
from dataclasses import dataclass
from pathlib import Path
from tkinter import filedialog, messagebox, ttk
from typing import Any, cast

from PIL import Image, ImageDraw, ImageTk
from tkinterdnd2 import DND_FILES, TkinterDnD
from matplotlib import get_data_path
from matplotlib.axes import Axes
from matplotlib.backend_bases import MouseButton, MouseEvent
from matplotlib.backends._backend_tk import NavigationToolbar2Tk
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
from matplotlib.figure import Figure
from matplotlib.lines import Line2D
from matplotlib.text import Annotation
from matplotlib.widgets import RectangleSelector
from openpyxl import load_workbook


WINDOW_BG = "#101715"
PANEL_BG = "#18221f"
RAISED_BG = "#22302b"
BORDER = "#34443e"
TEXT = "#dbe7e1"
MUTED_TEXT = "#91a69d"
ACCENT = "#48d6ad"


@dataclass
class Signal:
    uid: str
    source: str
    name: str
    times: list[float]
    values: list[float]

    @property
    def label(self) -> str:
        return f"{Path(self.source).stem} / {self.name}"


def _number(value: object) -> float | None:
    if not isinstance(value, (str, int, float)) or isinstance(value, bool):
        return None
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    return result if math.isfinite(result) else None


def _signal_color(index: int) -> str:
    red, green, blue = colorsys.hsv_to_rgb((index * 0.618033988749895) % 1, 0.68, 0.94)
    return f"#{round(red * 255):02x}{round(green * 255):02x}{round(blue * 255):02x}"


def read_dump(path: str | Path) -> tuple[str, list[Signal]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        if sheet is None:
            raise ValueError("The workbook has no active worksheet.")
        header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if header is None or not header:
            raise ValueError("The selected worksheet is empty.")

        headers = [str(value).strip() if value is not None else "" for value in header]
        time_index = next(
            (index for index, name in enumerate(headers) if name.casefold() in {"time", "time [s]", "time(s)"}),
            0,
        )
        signal_columns = [
            (index, name)
            for index, name in enumerate(headers)
            if name and index != time_index
        ]
        times: list[float] = []
        values = {index: [] for index, _ in signal_columns}

        for row in sheet.iter_rows(min_row=2, values_only=True):
            time_value = _number(row[time_index] if time_index < len(row) else None)
            if time_value is None:
                continue
            times.append(time_value)
            for index, _ in signal_columns:
                value = _number(row[index] if index < len(row) else None)
                values[index].append(value if value is not None else math.nan)

        if not times:
            raise ValueError("No numeric time samples were found in the worksheet.")

        signals = []
        source = str(Path(path).resolve())
        for index, name in signal_columns:
            samples = values[index]
            if any(math.isfinite(value) for value in samples):
                signals.append(
                    Signal(
                        uid=f"{source}:{index}",
                        source=source,
                        name=name,
                        times=times,
                        values=samples,
                    )
                )
        return sheet.title, signals
    finally:
        workbook.close()


class DaqIdeaPlotter:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.root.title("DaqIdea Graph Panel")
        self.root.geometry("1440x900")
        self.root.minsize(1060, 680)
        self.root.configure(bg=WINDOW_BG)

        self.datasets: list[tuple[str, str, list[Signal]]] = []
        self.signals: dict[str, Signal] = {}
        self.plotted: list[tuple[str, str]] = []
        self.signal_visibility: dict[str, tk.BooleanVar] = {}
        self.signal_colors: dict[str, str] = {}
        self.signal_tag_by_uid: dict[str, str] = {}
        self.signal_tree: ttk.Treeview | None = None
        self.display_cache: dict[str, tuple[int, list[float], list[float]]] = {}
        self.executor = ThreadPoolExecutor(max_workers=2, thread_name_prefix="daqidea-loader")
        self.load_futures: dict[Future[tuple[str, list[Signal]]], str] = {}
        self.loading_paths: set[str] = set()
        self.load_poll_id: str | None = None
        self.resize_job_id: str | None = None
        self.toolbar_icons: dict[str, ImageTk.PhotoImage] = {}
        self.separate_graphs_mode = False
        self.samples_visible = False
        self.plot_axes: list[Axes] = []
        self.right_plot_axes: list[Axes] = []
        self.signal_axes: dict[str, Axes] = {}
        self.rename_entry: tk.Entry | None = None
        self.rename_uid: str | None = None
        self.horizontal_scroll_remainder = 0.0
        self.drag_source_uid: str | None = None
        self.drag_target_uid: str | None = None
        self.shortcuts_window: tk.Toplevel | None = None
        self.area_selector: RectangleSelector | None = None
        self.zoom_mode = False
        self.zoom_axis_mode = "both"
        self.cursor_seconds: dict[str, float] = {}
        self.cursor_mode: str | None = None
        self.cursor_artists: dict[str, list[Line2D]] = {}
        self.dragging_cursor: str | None = None
        self.live_cursor_mode = False
        self.live_cursor_seconds: float | None = None
        self.live_cursor_artist: Line2D | None = None
        self.live_cursor_artists: list[Line2D] = []
        self.live_popup_artist: Annotation | None = None
        self.live_popup_uid: str | None = None
        self.measurement_refresh_id: str | None = None
        self.time_unit = tk.StringVar(value="s")
        self.live_cursor_text = tk.StringVar(value="t: --")
        self.cursor_delta_text = tk.StringVar(value="Δt: --")
        self.status_text = tk.StringVar(value="Open an Excel dump to begin.")

        self._style()
        self._load_toolbar_icons()
        self._build_ui()
        self.root.protocol("WM_DELETE_WINDOW", self._close)
        self._draw_plot(reset_view=True)

    def _style(self) -> None:
        style = ttk.Style(self.root)
        style.theme_use("clam")
        style.configure("TFrame", background=WINDOW_BG)
        style.configure("Panel.TFrame", background=PANEL_BG)
        style.configure("TLabel", background=WINDOW_BG, foreground=TEXT, font=("Segoe UI", 9))
        style.configure("Title.TLabel", background=WINDOW_BG, foreground=TEXT, font=("Segoe UI Semibold", 17))
        style.configure("Section.TLabel", background=PANEL_BG, foreground=MUTED_TEXT, font=("Segoe UI Semibold", 9))
        style.configure("Status.TLabel", background=RAISED_BG, foreground=MUTED_TEXT, font=("Segoe UI", 9))
        style.configure("TButton", font=("Segoe UI Semibold", 9), padding=(10, 7), background=RAISED_BG, foreground=TEXT, bordercolor=BORDER)
        style.map("TButton", background=[("active", "#2d3c36"), ("pressed", "#34463f")], foreground=[("disabled", "#64736c")])
        style.configure("Accent.TButton", background="#1b735c", foreground="#f1fff9", bordercolor="#29866b")
        style.map("Accent.TButton", background=[("active", "#25866b"), ("pressed", "#17634f")])
        style.configure("Active.TButton", background="#8b702d", foreground="#fff8e5", bordercolor="#b29445")
        style.map("Active.TButton", background=[("active", "#a28337")])
        style.configure("Icon.TButton", font=("Segoe UI Symbol", 11), padding=(3, 3), background=RAISED_BG, foreground=TEXT, bordercolor=BORDER, width=3)
        style.map("Icon.TButton", background=[("active", "#2d3c36")])
        style.configure("Icon.TMenubutton", font=("Segoe UI Symbol", 11), padding=(3, 3), background=RAISED_BG, foreground=TEXT, bordercolor=BORDER, width=3)
        style.map("Icon.TMenubutton", background=[("active", "#2d3c36")])
        style.configure("Active.TMenubutton", font=("Segoe UI Symbol", 11), padding=(3, 3), background="#8b702d", foreground="#fff8e5", bordercolor="#b29445", width=3)
        style.map("Active.TMenubutton", background=[("active", "#a28337")])
        style.configure("Treeview", background=PANEL_BG, fieldbackground=PANEL_BG, foreground=TEXT, rowheight=21, font=("Segoe UI", 9), bordercolor=BORDER)
        style.configure("Treeview.Heading", background=RAISED_BG, foreground=TEXT, font=("Segoe UI Semibold", 9), relief="flat")
        style.map("Treeview", background=[("selected", "#245b49")], foreground=[("selected", "#f1fff9")])
        style.configure("TCheckbutton", background=PANEL_BG, foreground=TEXT, font=("Segoe UI", 9))
        style.map("TCheckbutton", background=[("active", PANEL_BG)], foreground=[("disabled", "#64736c")])
        style.configure("TEntry", padding=6, fieldbackground=RAISED_BG, foreground=TEXT, insertcolor=TEXT)
        style.configure("Rename.TEntry", padding=3, fieldbackground="#f1f5f3", foreground="#101715", insertcolor="#101715", selectbackground="#27765e", selectforeground="#ffffff")
        style.configure("TCombobox", padding=5, fieldbackground=RAISED_BG, background=RAISED_BG, foreground=TEXT, arrowcolor=TEXT)
        style.map("TCombobox", fieldbackground=[("readonly", RAISED_BG)], foreground=[("readonly", TEXT)], selectbackground=[("readonly", RAISED_BG)], selectforeground=[("readonly", TEXT)])
        style.configure("TScrollbar", background=RAISED_BG, troughcolor=PANEL_BG, bordercolor=PANEL_BG, arrowcolor=MUTED_TEXT)
        style.configure("TSeparator", background=BORDER)

    def _load_toolbar_icons(self) -> None:
        image_dir = Path(get_data_path()) / "images"
        icon_files = (
            ("home", "home.png"),
            ("back", "back.png"),
            ("forward", "forward.png"),
            ("pan", "move.png"),
            ("zoom", "zoom_to_rect.png"),
            ("subplots", "subplots.png"),
            ("cursor", "move.png"),
            ("save", "filesave.png"),
        )
        for name, filename in icon_files:
            with Image.open(image_dir / filename) as source:
                icon = source.convert("RGBA")
            tinted = Image.new("RGBA", icon.size, TEXT)
            tinted.putalpha(icon.getchannel("A"))
            self.toolbar_icons[name] = ImageTk.PhotoImage(
                tinted.resize((18, 18), Image.Resampling.LANCZOS),
                master=self.root,
            )

        export_icon = Image.new("RGBA", (24, 24), (0, 0, 0, 0))
        draw = ImageDraw.Draw(export_icon)
        draw.line((12, 3, 12, 15), fill=TEXT, width=2)
        draw.line((7, 10, 12, 15, 17, 10), fill=TEXT, width=2)
        draw.line((4, 17, 4, 21, 20, 21, 20, 17), fill=TEXT, width=2)
        self.toolbar_icons["export"] = ImageTk.PhotoImage(
            export_icon.resize((18, 18), Image.Resampling.LANCZOS),
            master=self.root,
        )

        samples_icon = Image.new("RGBA", (24, 24), (0, 0, 0, 0))
        draw = ImageDraw.Draw(samples_icon)
        draw.line((3, 19, 8, 13, 13, 16, 20, 6), fill=TEXT, width=1)
        for x, y in ((3, 19), (8, 13), (13, 16), (20, 6)):
            draw.ellipse((x - 2, y - 2, x + 2, y + 2), fill=ACCENT)
        self.toolbar_icons["samples"] = ImageTk.PhotoImage(
            samples_icon.resize((18, 18), Image.Resampling.LANCZOS),
            master=self.root,
        )

        info_icon = Image.new("RGBA", (24, 24), (0, 0, 0, 0))
        draw = ImageDraw.Draw(info_icon)
        draw.ellipse((3, 3, 21, 21), outline=TEXT, width=2)
        draw.ellipse((10, 7, 14, 11), fill=ACCENT)
        draw.line((12, 12, 12, 18), fill=ACCENT, width=2)
        self.toolbar_icons["info"] = ImageTk.PhotoImage(
            info_icon.resize((18, 18), Image.Resampling.LANCZOS),
            master=self.root,
        )

        info_icon = Image.new("RGBA", (24, 24), (0, 0, 0, 0))
        draw = ImageDraw.Draw(info_icon)
        draw.ellipse((3, 3, 21, 21), outline=TEXT, width=2)
        draw.ellipse((10, 7, 14, 11), fill=ACCENT)
        draw.line((12, 12, 12, 18), fill=ACCENT, width=2)
        self.toolbar_icons["info"] = ImageTk.PhotoImage(
            info_icon.resize((18, 18), Image.Resampling.LANCZOS),
            master=self.root,
        )

        open_icon = Image.new("RGBA", (24, 24), (0, 0, 0, 0))
        draw = ImageDraw.Draw(open_icon)
        draw.line((3, 8, 9, 8, 11, 10, 21, 10, 20, 20, 4, 20, 3, 8), fill=TEXT, width=2)
        draw.line((4, 11, 20, 11), fill=TEXT, width=2)
        draw.line((12, 15, 12, 8, 9, 11), fill=TEXT, width=2)
        draw.line((12, 8, 15, 11), fill=TEXT, width=2)
        self.toolbar_icons["open"] = ImageTk.PhotoImage(open_icon.resize((18, 18), Image.Resampling.LANCZOS), master=self.root)

        for name, hide in (("select_all", False), ("unselect_all", True)):
            selection_icon = Image.new("RGBA", (24, 24), (0, 0, 0, 0))
            draw = ImageDraw.Draw(selection_icon)
            for y in (5, 11, 17):
                draw.line((8, y, 21, y), fill=TEXT, width=2)
            if hide:
                draw.line((2, 8, 6, 12, 2, 16), fill=TEXT, width=2)
            else:
                draw.line((2, 11, 4, 13, 7, 8), fill=TEXT, width=2)
            self.toolbar_icons[name] = ImageTk.PhotoImage(
                selection_icon.resize((18, 18), Image.Resampling.LANCZOS),
                master=self.root,
            )

    def _add_tooltip(self, widget: tk.Misc, text: str) -> None:
        pending_after: list[str | None] = [None]
        tooltip_window: list[tk.Toplevel | None] = [None]

        def hide(_event: tk.Event | None = None) -> None:
            after_id = pending_after[0]
            if after_id is not None:
                self.root.after_cancel(after_id)
                pending_after[0] = None
            window = tooltip_window[0]
            if window is not None:
                window.destroy()
                tooltip_window[0] = None

        def show() -> None:
            pending_after[0] = None
            if not widget.winfo_exists():
                return
            window = tk.Toplevel(self.root)
            window.wm_overrideredirect(True)
            window.configure(background=RAISED_BG)
            tk.Label(window, text=text, background=RAISED_BG, foreground=TEXT, padx=8, pady=4, font=("Segoe UI", 9)).pack()
            window.update_idletasks()
            tooltip_width = window.winfo_reqwidth()
            tooltip_height = window.winfo_reqheight()
            root_top = self.root.winfo_rooty()
            pointer_x = widget.winfo_pointerx()
            pointer_y = widget.winfo_pointery()
            above_y = pointer_y - tooltip_height - 8
            below_y = pointer_y + 14
            y = above_y if above_y >= root_top + 4 else below_y
            x = max(4, min(pointer_x + 10, window.winfo_screenwidth() - tooltip_width - 4))
            y = max(4, min(y, window.winfo_screenheight() - tooltip_height - 4))
            window.wm_geometry(f"+{x}+{y}")
            tooltip_window[0] = window

        def schedule(_event: tk.Event) -> None:
            hide()
            pending_after[0] = self.root.after(450, show)

        widget.bind("<Enter>", schedule, add="+")
        widget.bind("<Leave>", hide, add="+")
        widget.bind("<ButtonPress>", hide, add="+")

    def _add_menu_tooltips(self, menu: tk.Menu, descriptions: dict[str, str]) -> None:
        pending_after: list[str | None] = [None]
        tooltip_window: list[tk.Toplevel | None] = [None]
        active_description: list[str | None] = [None]

        def hide(_event: tk.Event | None = None) -> None:
            if pending_after[0] is not None:
                self.root.after_cancel(pending_after[0])
                pending_after[0] = None
            if tooltip_window[0] is not None:
                tooltip_window[0].destroy()
                tooltip_window[0] = None
            active_description[0] = None

        def show(description: str) -> None:
            pending_after[0] = None
            if not menu.winfo_exists():
                return
            window = tk.Toplevel(self.root)
            window.wm_overrideredirect(True)
            window.configure(background=RAISED_BG)
            window.wm_geometry(f"+{menu.winfo_pointerx() + 12}+{menu.winfo_pointery() + 16}")
            tk.Label(window, text=description, background=RAISED_BG, foreground=TEXT, padx=8, pady=4, font=("Segoe UI", 9)).pack()
            tooltip_window[0] = window

        def on_motion(_event: tk.Event) -> None:
            active_index = menu.index("active")
            if not isinstance(active_index, int):
                hide()
                return
            label = str(menu.entrycget(active_index, "label"))
            description = descriptions.get(label)
            if description == active_description[0]:
                return
            hide()
            if description:
                active_description[0] = description
                pending_after[0] = self.root.after(450, lambda: show(description))

        menu.bind("<Motion>", on_motion, add="+")
        menu.bind("<Leave>", hide, add="+")
        menu.bind("<Unmap>", hide, add="+")

    def _build_ui(self) -> None:
        body = ttk.Panedwindow(self.root, orient="horizontal")
        body.pack(fill="both", expand=True, padx=6, pady=(0, 3))

        sidebar = ttk.Frame(body, style="Panel.TFrame", padding=12, width=365)
        sidebar.pack_propagate(False)
        body.add(sidebar, weight=0)
        ttk.Label(sidebar, text="SIGNALS & VALUES", style="Section.TLabel").pack(anchor="w", pady=(0, 8))
        signal_list = ttk.Frame(sidebar, style="Panel.TFrame")
        signal_list.pack(fill="both", expand=True)
        self.signal_tree = ttk.Treeview(
            signal_list,
            columns=("axis", "show", "live", "cursor_a", "cursor_b", "delta"),
        )
        self.signal_tree.configure(show="tree headings", selectmode="extended")
        self.signal_tree.heading("#0", text="Signal")
        self.signal_tree.heading("axis", text="Y")
        self.signal_tree.heading("show", text="On")
        self.signal_tree.heading("live", text="Now")
        self.signal_tree.heading("cursor_a", text="A")
        self.signal_tree.heading("cursor_b", text="B")
        self.signal_tree.heading("delta", text="Δ")
        self.signal_tree.column("#0", width=210, minwidth=96, stretch=False)
        self.signal_tree.column("axis", width=32, minwidth=30, stretch=False, anchor="center")
        self.signal_tree.column("show", width=35, minwidth=34, stretch=False, anchor="center")
        self.signal_tree.column("live", width=55, minwidth=50, stretch=False, anchor="e")
        for column in ("cursor_a", "cursor_b"):
            self.signal_tree.column(column, width=48, minwidth=44, stretch=False, anchor="e")
        self.signal_tree.column("delta", width=52, minwidth=48, stretch=False, anchor="e")
        signal_scroll = ttk.Scrollbar(signal_list, orient="vertical", command=self.signal_tree.yview)
        horizontal_scroll = ttk.Scrollbar(signal_list, orient="horizontal", command=self.signal_tree.xview)
        self.signal_tree.configure(yscrollcommand=signal_scroll.set, xscrollcommand=horizontal_scroll.set)
        self.signal_tree.grid(row=0, column=0, sticky="nsew")
        signal_scroll.grid(row=0, column=1, sticky="ns")
        horizontal_scroll.grid(row=1, column=0, sticky="ew")
        signal_list.rowconfigure(0, weight=1)
        signal_list.columnconfigure(0, weight=1)
        self.signal_tree.bind("<Button-1>", self._on_signal_list_click)
        self.signal_tree.bind("<<TreeviewSelect>>", self._on_signal_selection_changed, add="+")
        self.signal_tree.bind("<Double-1>", self._begin_signal_rename)
        self.signal_tree.bind("<B1-Motion>", self._on_signal_drag_motion)
        self.signal_tree.bind("<ButtonRelease-1>", self._on_signal_drag_release)
        self.signal_tree.bind("<Shift-MouseWheel>", self._scroll_signal_table_horizontally)
        self._add_tooltip(self.signal_tree, "Ctrl/Shift-click signals to multi-select. Drag rows to reorder within a dump. Double-click names to edit.")
        list_buttons = ttk.Frame(sidebar, style="Panel.TFrame")
        list_buttons.pack(fill="x", pady=(8, 0))
        select_all_button = ttk.Button(list_buttons, image=self.toolbar_icons["select_all"], style="Icon.TButton", command=lambda: self._set_selected_visible(True))
        select_all_button.pack(side="left", padx=(0, 4))
        self._add_tooltip(select_all_button, "Show selected signals, or all signals if none are selected")
        unselect_all_button = ttk.Button(list_buttons, image=self.toolbar_icons["unselect_all"], style="Icon.TButton", command=lambda: self._set_selected_visible(False))
        unselect_all_button.pack(side="left")
        self._add_tooltip(unselect_all_button, "Hide selected signals, or all signals if none are selected")
        open_button = ttk.Button(
            sidebar,
            text="Open dump",
            image=self.toolbar_icons["open"],
            compound="left",
            style="Accent.TButton",
            command=self._open_dump,
        )
        open_button.pack(side="bottom", fill="x", pady=(8, 0))
        self._add_tooltip(open_button, "Open one or more DaqIdea Excel workbooks")

        graph_panel = ttk.Frame(body, style="Panel.TFrame", padding=(4, 4, 4, 3))
        body.add(graph_panel, weight=1)
        graph_toolbar = ttk.Frame(graph_panel, style="Panel.TFrame")
        graph_toolbar.pack(fill="x", pady=(0, 3))
        ttk.Label(graph_toolbar, text="TIME", style="Section.TLabel").pack(side="left", padx=(0, 6))
        time_unit_combo = ttk.Combobox(graph_toolbar, textvariable=self.time_unit, values=("s", "ms"), width=5, state="readonly")
        time_unit_combo.pack(side="left")
        self._add_tooltip(time_unit_combo, "Set the graph time axis to seconds or milliseconds")
        self.time_unit.trace_add("write", self._on_time_unit_changed)
        time_unit_combo.bind("<KeyPress>", self._on_keyboard_shortcut, add="+")
        ttk.Separator(graph_toolbar, orient="vertical").pack(side="left", fill="y", padx=11)
        self.zoom_menu = tk.Menu(self.root, tearoff=0, background=RAISED_BG, foreground=TEXT, activebackground="#2d3c36", activeforeground=TEXT, borderwidth=1)
        self.zoom_menu.add_command(label="Area zoom", command=lambda: self._toggle_area_zoom("both"))
        self.zoom_menu.add_command(label="X-axis area zoom", command=lambda: self._toggle_area_zoom("x"))
        self.zoom_menu.add_command(label="Y-axis area zoom", command=lambda: self._toggle_area_zoom("y"))
        self.zoom_menu.add_separator()
        self.zoom_menu.add_command(label="Fit all axes", command=self._fit_all)
        self.zoom_menu.add_command(label="Fit X axis", command=self._fit_x)
        self.zoom_menu.add_command(label="Fit Y axes", command=self._fit_y)
        self._add_menu_tooltips(
            self.zoom_menu,
            {
                "Area zoom": "Drag a rectangle to zoom both axes",
                "X-axis area zoom": "Drag a range to zoom only the time axis",
                "Y-axis area zoom": "Drag a range to zoom only the value axis",
                "Fit all axes": "Fit X and Y to the visible signals",
                "Fit X axis": "Fit time while preserving the current Y range",
                "Fit Y axes": "Fit value axes while preserving the current time range",
            },
        )
        self.zoom_menu_button = ttk.Menubutton(
            graph_toolbar,
            image=self.toolbar_icons["zoom"],
            menu=self.zoom_menu,
            direction="below",
            style="Icon.TMenubutton",
        )
        self.zoom_menu_button.pack(side="left", padx=(0, 5))
        self.zoom_button = self.zoom_menu_button
        self._add_tooltip(self.zoom_menu_button, "Zoom and fit options; Shift+wheel zooms X, Ctrl+wheel zooms Y")
        ttk.Separator(graph_toolbar, orient="vertical").pack(side="left", fill="y", padx=5)
        self.cursor_menu = tk.Menu(self.root, tearoff=0, background=RAISED_BG, foreground=TEXT, activebackground="#2d3c36", activeforeground=TEXT, borderwidth=1)
        self.cursor_menu.add_command(label="Place cursor A", command=lambda: self._arm_cursor("A"))
        self.cursor_menu.add_command(label="Place cursor B", command=lambda: self._arm_cursor("B"))
        self.cursor_menu.add_command(label="Toggle live cursor", command=self._toggle_live_cursor)
        self.cursor_menu.add_separator()
        self.cursor_menu.add_command(label="Clear cursors", command=self._clear_cursors)
        self._add_menu_tooltips(
            self.cursor_menu,
            {
                "Place cursor A": "Place or move cursor A on the graph",
                "Place cursor B": "Place or move cursor B on the graph",
                "Toggle live cursor": "Show a crosshair and live values under the pointer",
                "Clear cursors": "Remove A/B cursor lines and measurements",
            },
        )
        self.cursor_menu_button = ttk.Menubutton(
            graph_toolbar,
            image=self.toolbar_icons["cursor"],
            menu=self.cursor_menu,
            direction="below",
            style="Icon.TMenubutton",
        )
        self.cursor_menu_button.pack(side="left", padx=(0, 5))
        self._add_tooltip(self.cursor_menu_button, "Cursor placement, live cursor, and clear")
        info_button = ttk.Button(graph_toolbar, image=self.toolbar_icons["info"], style="Icon.TButton", command=self._show_shortcuts)
        info_button.pack(side="left", padx=(0, 5))
        self._add_tooltip(info_button, "Keyboard shortcuts and drag/drop help")
        self.separate_graphs_button = ttk.Button(
            graph_toolbar,
            image=self.toolbar_icons["subplots"],
            style="Icon.TButton",
            command=self._toggle_separate_graphs,
        )
        self.separate_graphs_button.pack(side="left", padx=(0, 5))
        self._add_tooltip(self.separate_graphs_button, "Stack visible signals in separate plots with independent Y ranges")
        self.samples_button = ttk.Button(
            graph_toolbar,
            image=self.toolbar_icons["samples"],
            style="Icon.TButton",
            command=self._toggle_samples,
        )
        self.samples_button.pack(side="left", padx=(0, 5))
        self._add_tooltip(self.samples_button, "Toggle a dot at every raw sample for visible signals")
        self.file_menu = tk.Menu(self.root, tearoff=0, background=RAISED_BG, foreground=TEXT, activebackground="#2d3c36", activeforeground=TEXT, borderwidth=1)
        self.file_menu.add_command(label="Save graph image", command=self._save_graph)
        self.file_menu.add_command(label="Export visible signals to CSV", command=self._export_csv)
        self._add_menu_tooltips(
            self.file_menu,
            {
                "Save graph image": "Save the current graph as PNG, PDF, or SVG",
                "Export visible signals to CSV": "Export original samples for visible signals",
            },
        )
        file_menu_button = ttk.Menubutton(
            graph_toolbar,
            image=self.toolbar_icons["save"],
            menu=self.file_menu,
            direction="below",
            style="Icon.TMenubutton",
        )
        file_menu_button.pack(side="right")
        self._add_tooltip(file_menu_button, "Save and export actions")

        self.figure = Figure(figsize=(9, 5), dpi=100, facecolor="#111a19")
        self.ax = self.figure.add_subplot(111)
        self.ax_right = self.ax.twinx()
        self.plot_axes = [self.ax]
        self.right_plot_axes = [self.ax_right]
        self.canvas = FigureCanvasTkAgg(self.figure, master=graph_panel)
        self.canvas_widget = self.canvas.get_tk_widget()
        self.canvas_widget.configure(highlightthickness=0)
        self.canvas_widget.pack(fill="both", expand=True)
        for target in (self.root, self.signal_tree, self.canvas_widget):
            dnd_target = cast(Any, target)
            register_drop = getattr(dnd_target, "drop_target_register", None)
            bind_drop = getattr(dnd_target, "dnd_bind", None)
            if callable(register_drop) and callable(bind_drop):
                register_drop(DND_FILES)
                bind_drop("<<Drop>>", self._on_file_drop)
        self._add_tooltip(self.canvas_widget, "Scroll to zoom; Shift+scroll zooms X, Ctrl+scroll zooms Y. Drag A/B cursor lines to move them.")
        self.canvas.mpl_connect("button_press_event", self._on_plot_click)
        self.canvas.mpl_connect("button_release_event", self._on_plot_release)
        self.canvas.mpl_connect("motion_notify_event", self._on_plot_motion)
        self.canvas.mpl_connect("resize_event", self._on_canvas_resize)
        self.canvas.mpl_connect("scroll_event", self._on_scroll_zoom)
        self.area_selector = RectangleSelector(
            self.ax,
            self._on_area_zoom,
            useblit=True,
            button=[MouseButton.LEFT],
            minspanx=2,
            minspany=2,
            spancoords="pixels",
            props={"facecolor": "#48d6ad", "edgecolor": "#a5f3d6", "alpha": 0.18},
        )
        self.area_selector.set_active(False)
        self.nav_toolbar = NavigationToolbar2Tk(self.canvas, graph_panel, pack_toolbar=False)
        self.nav_toolbar.update()
        self.nav_toolbar.configure(background=PANEL_BG)
        toolbar_icons = {
            "Home": "home",
            "Back": "back",
            "Forward": "forward",
            "Pan": "pan",
            "Zoom": "zoom",
            "Subplots": "subplots",
            "Save": "save",
        }
        toolbar_descriptions = {
            label: description
            for label, description, _image, _method in self.nav_toolbar.toolitems
            if label and description
        }
        for toolbar_widget in self.nav_toolbar.winfo_children():
            if isinstance(toolbar_widget, tk.Button):
                label = toolbar_widget.cget("text")
                toolbar_widget.configure(
                    background=RAISED_BG,
                    foreground=TEXT,
                    activebackground="#2d3c36",
                    activeforeground=TEXT,
                    relief="flat",
                    borderwidth=0,
                    highlightthickness=0,
                )
                if label in toolbar_icons:
                    toolbar_widget.configure(image=self.toolbar_icons[toolbar_icons[label]], text="", compound="center")
                self._add_tooltip(toolbar_widget, toolbar_descriptions.get(label) or str(label or "Navigation tool"))
            elif isinstance(toolbar_widget, tk.Checkbutton):
                label = toolbar_widget.cget("text")
                toolbar_widget.configure(
                    background=RAISED_BG,
                    foreground=TEXT,
                    activebackground="#2d3c36",
                    activeforeground=TEXT,
                    selectcolor=RAISED_BG,
                    relief="flat",
                    borderwidth=0,
                    highlightthickness=0,
                )
                if label in toolbar_icons:
                    toolbar_widget.configure(image=self.toolbar_icons[toolbar_icons[label]], text="", compound="center")
                self._add_tooltip(toolbar_widget, toolbar_descriptions.get(label) or str(label or "Navigation tool"))
            elif isinstance(toolbar_widget, tk.Label):
                toolbar_widget.configure(background=PANEL_BG, foreground=MUTED_TEXT)
        self.nav_toolbar.pack(fill="x", pady=(2, 6))
        self.nav_toolbar._message_label.pack_forget()

        status_bar = ttk.Frame(self.root, style="Status.TLabel")
        status_bar.pack(fill="x")
        status = ttk.Label(status_bar, textvariable=self.status_text, anchor="w", padding=(12, 6), style="Status.TLabel")
        status.pack(side="left", fill="x", expand=True)
        ttk.Label(status_bar, textvariable=self.cursor_delta_text, anchor="e", padding=(8, 6), style="Status.TLabel").pack(side="right")
        ttk.Label(status_bar, textvariable=self.live_cursor_text, anchor="e", padding=(8, 6), style="Status.TLabel").pack(side="right")
        self.root.bind_all("<KeyPress>", self._on_keyboard_shortcut, add="+")

    def _open_dump(self) -> None:
        paths = filedialog.askopenfilenames(
            title="Open DaqIdea Excel dump",
            filetypes=(("Excel workbooks", "*.xlsx *.xlsm"), ("All files", "*.*")),
        )
        if not paths:
            return
        self._load_paths(paths)

    def _on_file_drop(self, event: Any) -> str:
        try:
            paths = [str(Path(path)) for path in self.root.tk.splitlist(event.data)]
        except (TypeError, tk.TclError):
            paths = [str(Path(event.data))]
        workbooks = [path for path in paths if Path(path).suffix.casefold() in {".xlsx", ".xlsm"}]
        if workbooks:
            self._load_paths(workbooks)
        return "break"

    def _load_paths(self, paths: tuple[str, ...] | list[str]) -> None:
        scheduled = 0
        for path in paths:
            resolved_path = str(Path(path).resolve())
            if resolved_path in self.loading_paths or any(existing[0] == resolved_path for existing in self.datasets):
                continue
            self.loading_paths.add(resolved_path)
            future = self.executor.submit(read_dump, resolved_path)
            self.load_futures[future] = resolved_path
            scheduled += 1
        if scheduled:
            self.status_text.set(f"Loading {len(self.load_futures)} workbook(s) in the background...")
            self._schedule_load_poll()

    def _schedule_load_poll(self) -> None:
        if self.load_poll_id is None and self.load_futures:
            self.load_poll_id = self.root.after(50, self._poll_loads)

    def _poll_loads(self) -> None:
        self.load_poll_id = None
        loaded = 0
        errors = []
        for future, path in list(self.load_futures.items()):
            if not future.done():
                continue
            del self.load_futures[future]
            self.loading_paths.discard(path)
            try:
                sheet_name, signals = future.result()
                self.datasets.append((path, sheet_name, signals))
                for signal in signals:
                    self.signals[signal.uid] = signal
                    self.signal_visibility[signal.uid] = tk.BooleanVar(master=self.root, value=True)
                    self.signal_colors[signal.uid] = _signal_color(len(self.signal_colors))
                    self.plotted.append((signal.uid, "left"))
                loaded += 1
            except Exception as error:
                errors.append(f"{Path(path).name}: {error}")

        if loaded:
            self._refresh_signal_list()
            self._draw_plot(reset_view=True)
            visible_count = sum(variable.get() for variable in self.signal_visibility.values())
            self.status_text.set(f"Loaded {loaded} dump{'s' if loaded != 1 else ''}; plotting {visible_count} signals.")
        elif self.load_futures:
            self.status_text.set(f"Loading {len(self.load_futures)} workbook(s) in the background...")
        if errors:
            messagebox.showerror("Could not open dump", "\n".join(errors), parent=self.root)
        self._schedule_load_poll()

    def _close(self) -> None:
        if self.load_poll_id is not None:
            self.root.after_cancel(self.load_poll_id)
        if self.resize_job_id is not None:
            self.root.after_cancel(self.resize_job_id)
        if self.measurement_refresh_id is not None:
            self.root.after_cancel(self.measurement_refresh_id)
        self.executor.shutdown(wait=False, cancel_futures=True)
        self.root.destroy()

    def _refresh_signal_list(self) -> None:
        tree = self.signal_tree
        if tree is None:
            return
        tree.delete(*tree.get_children())
        grouped = True
        axis_by_uid: dict[str, str] = {uid: axis for uid, axis in self.plotted}
        color_index = 0
        tree.tag_configure("dataset-heading", foreground=MUTED_TEXT, font=("Segoe UI Semibold", 9))
        for dataset_index, (path, _sheet_name, signals) in enumerate(self.datasets):
            parent = f"dataset:{dataset_index}"
            all_visible = bool(signals) and all(self.signal_visibility[signal.uid].get() for signal in signals)
            tree.insert(
                "",
                "end",
                iid=parent,
                text=Path(path).name,
                values=("", "✓" if all_visible else "", "", "", "", ""),
                tags=("dataset-heading",),
                open=True,
            )
            for signal in signals:
                uid = signal.uid
                axis = axis_by_uid.get(uid, "left")
                tag = f"signal-color-{color_index}"
                self.signal_tag_by_uid[uid] = tag
                tree.tag_configure(tag, foreground=self.signal_colors[uid])
                tree.insert(
                    parent,
                    "end",
                    iid=uid,
                    text=f"●  {signal.name}",
                    values=("L" if axis == "left" else "R", "✓" if self.signal_visibility[uid].get() else "", "", "", "", ""),
                    tags=(tag,),
                )
                color_index += 1

    def _on_signal_list_click(self, event: tk.Event) -> str | None:
        tree = self.signal_tree
        if tree is None:
            return None
        uid = tree.identify_row(event.y)
        column = tree.identify_column(event.x)
        if uid.startswith("dataset:") and column == "#2":
            dataset_index = int(uid.split(":", 1)[1])
            if dataset_index < len(self.datasets):
                signals = self.datasets[dataset_index][2]
                all_visible = bool(signals) and all(self.signal_visibility[signal.uid].get() for signal in signals)
                self._set_dataset_visible(dataset_index, not all_visible)
            return "break"
        if uid not in self.signals:
            return None
        if column == "#0":
            self.drag_source_uid = uid
        if column == "#1":
            current_axis = next(axis for signal_uid, axis in self.plotted if signal_uid == uid)
            self._set_signal_axis(uid, "right" if current_axis == "left" else "left")
            return None
        if column == "#2":
            self.signal_visibility[uid].set(not self.signal_visibility[uid].get())
            self._visibility_changed(uid)
            return None
        return None

    def _on_signal_drag_motion(self, event: tk.Event) -> None:
        tree = self.signal_tree
        if tree is None or self.drag_source_uid not in self.signals:
            return
        target_uid = tree.identify_row(event.y)
        if target_uid in self.signals and tree.parent(target_uid) == tree.parent(self.drag_source_uid):
            tree.configure(cursor="fleur")
        else:
            tree.configure(cursor="no")

    def _on_signal_drag_release(self, event: tk.Event) -> None:
        tree = self.signal_tree
        source_uid = self.drag_source_uid
        if tree is None or source_uid not in self.signals:
            self.drag_source_uid = None
            return
        target_uid = tree.identify_row(event.y)
        tree.configure(cursor="")
        self.drag_source_uid = None
        if target_uid not in self.signals or target_uid == source_uid or tree.parent(target_uid) != tree.parent(source_uid):
            return
        parent = tree.parent(source_uid)
        dataset_index = int(parent.split(":", 1)[1])
        signals = self.datasets[dataset_index][2]
        source_index = next(index for index, signal in enumerate(signals) if signal.uid == source_uid)
        target_index = next(index for index, signal in enumerate(signals) if signal.uid == target_uid)
        moved_signal = signals.pop(source_index)
        signals.insert(target_index, moved_signal)
        plotted_index = next(index for index, (uid, _axis) in enumerate(self.plotted) if uid == source_uid)
        moved_plot = self.plotted.pop(plotted_index)
        target_plot_index = next(index for index, (uid, _axis) in enumerate(self.plotted) if uid == target_uid)
        self.plotted.insert(target_plot_index, moved_plot)
        self._refresh_signal_list()
        self._draw_plot(reset_view=False)
        self.status_text.set(f"Moved {self.signals[source_uid].name}.")

    def _begin_signal_rename(self, event: tk.Event) -> str | None:
        tree = self.signal_tree
        if tree is None or tree.identify_column(event.x) != "#0":
            return None
        uid = tree.identify_row(event.y)
        if uid not in self.signals:
            return None
        bounds = tree.bbox(uid, "#0")
        if not bounds:
            return None
        x, y, width, height = bounds
        self._cancel_signal_rename()
        entry = tk.Entry(
            tree,
            background="#f1f5f3",
            foreground="#101715",
            insertbackground="#101715",
            selectbackground="#27765e",
            selectforeground="#ffffff",
            relief="solid",
            borderwidth=1,
            font=("Segoe UI", 9),
            justify="left",
        )
        entry.insert(0, self.signals[uid].name)
        edit_width = min(max(width, 320), max(100, tree.winfo_width() - x - 3))
        entry.place(x=x, y=y, width=edit_width, height=height)
        entry.lift()
        self.rename_entry = entry
        self.rename_uid = uid
        entry.bind("<Return>", self._commit_signal_rename)
        entry.bind("<Escape>", self._cancel_signal_rename)
        entry.bind("<FocusOut>", self._commit_signal_rename)
        entry.focus_force()
        entry.selection_range(0, "end")
        entry.xview_moveto(0)
        return "break"

    def _commit_signal_rename(self, _event: tk.Event | None = None) -> str:
        entry = self.rename_entry
        uid = self.rename_uid
        if entry is None or uid is None:
            return "break"
        new_name = entry.get().strip()
        entry.destroy()
        self.rename_entry = None
        self.rename_uid = None
        if new_name:
            signal = self.signals[uid]
            if new_name != signal.name:
                signal.name = new_name
                if self.signal_tree is not None and self.signal_tree.exists(uid):
                    self.signal_tree.item(uid, text=f"●  {new_name}")
                if self.separate_graphs_mode:
                    self._draw_plot(reset_view=False)
                self.status_text.set(f"Renamed signal to {new_name}.")
        return "break"

    def _cancel_signal_rename(self, _event: tk.Event | None = None) -> str:
        if self.rename_entry is not None:
            self.rename_entry.destroy()
        self.rename_entry = None
        self.rename_uid = None
        return "break"

    def _scroll_signal_table_horizontally(self, event: tk.Event) -> str:
        if self.signal_tree is None:
            return "break"
        self.horizontal_scroll_remainder += -event.delta / 120
        units = int(self.horizontal_scroll_remainder)
        if units:
            self.signal_tree.xview_scroll(units, "units")
            self.horizontal_scroll_remainder -= units
        return "break"

    def _on_signal_selection_changed(self, _event: tk.Event) -> None:
        self._refresh_live_popup()
        self.canvas.draw_idle()

    def _refresh_live_popup(self) -> None:
        tree = self.signal_tree
        selected = tree.selection() if tree is not None else ()
        uid = selected[0] if selected and selected[0] in self.signals else None
        if (
            not self.live_cursor_mode
            or self.live_cursor_seconds is None
            or uid is None
            or not self.signal_visibility[uid].get()
        ):
            self._hide_live_popup()
            return

        signal = self.signals[uid]
        sample_index = self._nearest_sample_index(signal, self.live_cursor_seconds)
        value = signal.values[sample_index]
        if not math.isfinite(value):
            self._hide_live_popup()
            return
        axis = self.signal_axes.get(uid, self.ax)
        scale = 1000 if self.time_unit.get() == "ms" else 1
        sample_time = signal.times[sample_index] * scale
        popup_text = f"{signal.name}\nt = {sample_time:.6g} {self.time_unit.get()}\nvalue = {value:.7g}"
        if self.live_popup_artist is not None and (
            self.live_popup_uid != uid or self.live_popup_artist.axes is not axis
        ):
            self._hide_live_popup()
        if self.live_popup_artist is None:
            self.live_popup_artist = axis.annotate(
                popup_text,
                xy=(sample_time, value),
                xytext=(14, 14),
                textcoords="offset points",
                color=TEXT,
                fontsize=8,
                ha="left",
                va="bottom",
                annotation_clip=True,
                bbox={
                    "boxstyle": "round,pad=0.35",
                    "facecolor": PANEL_BG,
                    "edgecolor": self.signal_colors[uid],
                    "alpha": 0.97,
                },
                arrowprops={"arrowstyle": "->", "color": self.signal_colors[uid], "lw": 0.9},
                zorder=10,
            )
            self.live_popup_uid = uid
        else:
            self.live_popup_artist.xy = (sample_time, value)
            self.live_popup_artist.set_text(popup_text)

    def _hide_live_popup(self) -> None:
        if self.live_popup_artist is not None:
            self.live_popup_artist.remove()
            self.live_popup_artist = None
        self.live_popup_uid = None

    def _update_signal_row(self, uid: str) -> None:
        tree = self.signal_tree
        if tree is None or not tree.exists(uid):
            return
        axis = next(axis for signal_uid, axis in self.plotted if signal_uid == uid)
        visible = self.signal_visibility[uid].get()
        current_values = tree.item(uid, "values")
        measurements = current_values[2:] if len(current_values) >= 6 else ("", "", "", "")
        tree.item(uid, values=("L" if axis == "left" else "R", "✓" if visible else "", *measurements))

    def _visibility_changed(self, uid: str) -> None:
        state = "shown" if self.signal_visibility[uid].get() else "hidden"
        self._update_signal_row(uid)
        self._draw_plot(reset_view=False)
        self.status_text.set(f"{self.signals[uid].name} {state}.")

    def _set_all_visible(self, visible: bool) -> None:
        for variable in self.signal_visibility.values():
            variable.set(visible)
        if self.signal_tree is not None:
            for uid, _axis in self.plotted:
                self._update_signal_row(uid)
        self._draw_plot(reset_view=False)
        self.status_text.set("All signals shown." if visible else "All signals hidden.")

    def _set_dataset_visible(self, dataset_index: int, visible: bool) -> None:
        if not 0 <= dataset_index < len(self.datasets):
            return
        for signal in self.datasets[dataset_index][2]:
            self.signal_visibility[signal.uid].set(visible)
            self._update_signal_row(signal.uid)
        if self.signal_tree is not None:
            parent = f"dataset:{dataset_index}"
            values = list(self.signal_tree.item(parent, "values"))
            values[1] = "✓" if visible else ""
            self.signal_tree.item(parent, values=values)
        self._draw_plot(reset_view=False)
        self.status_text.set(f"{Path(self.datasets[dataset_index][0]).name}: signals {'shown' if visible else 'hidden'}.")

    def _set_selected_visible(self, visible: bool) -> None:
        selected = [uid for uid in self.signal_tree.selection() if uid in self.signals] if self.signal_tree is not None else []
        if not selected:
            self._set_all_visible(visible)
            return
        for uid in selected:
            self.signal_visibility[uid].set(visible)
            self._update_signal_row(uid)
        self._draw_plot(reset_view=False)
        self.status_text.set(f"{len(selected)} selected signals {'shown' if visible else 'hidden'}.")

    def _show_shortcuts(self) -> None:
        if self.shortcuts_window is not None and self.shortcuts_window.winfo_exists():
            self.shortcuts_window.deiconify()
            self.shortcuts_window.lift()
            return
        window = tk.Toplevel(self.root)
        self.shortcuts_window = window
        window.title("Keyboard shortcuts")
        window.transient(self.root)
        window.resizable(False, False)
        window.configure(background=WINDOW_BG)
        content = ttk.Frame(window, padding=18, style="Panel.TFrame")
        content.pack(fill="both", expand=True)
        ttk.Label(content, text="KEYBOARD SHORTCUTS", style="Section.TLabel").grid(row=0, column=0, columnspan=2, sticky="w", pady=(0, 12))
        shortcuts = (
            ("Ctrl+O", "Open workbook"),
            ("Ctrl+S", "Save graph image"),
            ("Ctrl+E", "Export visible signals to CSV"),
            ("Ctrl+A", "Select all signals"),
            ("H / Shift+H", "Hide / show selected signals"),
            ("A / B", "Place cursor A / B"),
            ("L", "Toggle live cursor"),
            ("P", "Toggle sample dots"),
            ("G", "Toggle stacked plots"),
            ("Z", "Toggle area zoom mode"),
            ("1 / 2 / 3", "Fit all / X / Y axes"),
            ("I", "Show this shortcut list"),
            ("Esc", "Cancel current mode or dialog"),
        )
        for row, (key, action) in enumerate(shortcuts, start=1):
            ttk.Label(content, text=key, style="Section.TLabel").grid(row=row, column=0, sticky="w", padx=(0, 24), pady=3)
            ttk.Label(content, text=action, style="Body.TLabel").grid(row=row, column=1, sticky="w", pady=3)
        window.bind("<Escape>", lambda _event: window.destroy())
        window.protocol("WM_DELETE_WINDOW", window.destroy)
        window.bind("<Destroy>", lambda event: setattr(self, "shortcuts_window", None) if event.widget is window else None)

    def _on_keyboard_shortcut(self, event: tk.Event) -> str | None:
        keysym = str(event.keysym).lower()
        if keysym == "escape":
            if self.rename_entry is not None:
                return self._cancel_signal_rename()
            if self._dismiss_shortcuts():
                return "break"
            if self.zoom_mode:
                self._toggle_area_zoom(None)
            elif self.cursor_mode is not None:
                self.cursor_mode = None
                self._update_cursor_menu_style()
            elif self.live_cursor_mode:
                self._toggle_live_cursor()
            else:
                return None
            return "break"

        widget = event.widget
        read_only_combo = isinstance(widget, ttk.Combobox) and widget.instate(["readonly"])
        if isinstance(widget, (tk.Entry, ttk.Entry, tk.Text)) and not read_only_combo:
            return None
        if self.rename_entry is not None and widget is self.rename_entry:
            return None

        state = int(event.state)
        control = bool(state & 0x4) or bool(state & 0x8)
        shift = bool(state & 0x1)
        if control:
            if keysym not in {"o", "s", "e", "a"} or keysym == "a" and self.signal_tree is None:
                return None
            self._dismiss_shortcuts()
            if keysym == "o":
                self._open_dump()
            elif keysym == "s":
                self._save_graph()
            elif keysym == "e":
                self._export_csv()
            else:
                tree = self.signal_tree
                if tree is not None:
                    tree.selection_set(tuple(self.signals))
            return "break"

        actions = {
            "a": lambda: self._arm_cursor("A"),
            "b": lambda: self._arm_cursor("B"),
            "g": self._toggle_separate_graphs,
            "i": self._show_shortcuts,
            "l": self._toggle_live_cursor,
            "p": self._toggle_samples,
            "z": lambda: self._toggle_area_zoom("both"),
            "1": self._fit_all,
            "2": self._fit_x,
            "3": self._fit_y,
        }
        if keysym == "h":
            self._dismiss_shortcuts()
            self._set_selected_visible(shift)
        elif keysym in actions:
            if keysym != "i":
                self._dismiss_shortcuts()
            actions[keysym]()
        else:
            return None
        return "break"

    def _dismiss_shortcuts(self) -> bool:
        if self.shortcuts_window is None or not self.shortcuts_window.winfo_exists():
            return False
        self.shortcuts_window.destroy()
        return True

    def _set_signal_axis(self, uid: str, axis: str) -> None:
        self.plotted = [(signal_uid, axis if signal_uid == uid else current_axis) for signal_uid, current_axis in self.plotted]
        self._update_signal_row(uid)
        self._draw_plot(reset_view=False)

    def _visible_signals(self) -> list[tuple[str, str]]:
        return [(uid, axis) for uid, axis in self.plotted if self.signal_visibility[uid].get()]

    def _is_plot_axis(self, axis: object) -> bool:
        return axis in self.plot_axes or axis in self.right_plot_axes

    def _toggle_separate_graphs(self) -> None:
        self.separate_graphs_mode = not self.separate_graphs_mode
        self.separate_graphs_button.configure(style="Active.TButton" if self.separate_graphs_mode else "Icon.TButton")
        self._draw_plot(reset_view=True)
        self.status_text.set(
            "Showing each visible signal in its own diagram."
            if self.separate_graphs_mode
            else "Overlaying visible signals."
        )

    def _toggle_samples(self) -> None:
        self.samples_visible = not self.samples_visible
        self.samples_button.configure(style="Active.TButton" if self.samples_visible else "Icon.TButton")
        self._draw_plot(reset_view=False)
        self.status_text.set("Sample points shown." if self.samples_visible else "Sample points hidden.")

    def _update_cursor_menu_style(self) -> None:
        style = "Active.TMenubutton" if self.live_cursor_mode or self.cursor_mode is not None else "Icon.TMenubutton"
        self.cursor_menu_button.configure(style=style)

    def _toggle_area_zoom(self, axis_mode: str | None = None) -> None:
        if self.live_cursor_mode:
            self._deactivate_live_cursor()
        if self.zoom_mode and (axis_mode is None or axis_mode == self.zoom_axis_mode):
            self.zoom_mode = False
        else:
            self.zoom_mode = True
            self.zoom_axis_mode = axis_mode or "both"
        self.cursor_mode = None
        self._update_cursor_menu_style()
        self.zoom_button.configure(style="Active.TMenubutton" if self.zoom_mode else "Icon.TMenubutton")
        if self.area_selector is not None:
            self.area_selector.set_active(self.zoom_mode)
        if self.zoom_mode:
            mode_name = {"both": "both axes", "x": "X axis", "y": "Y axis"}[self.zoom_axis_mode]
            self.status_text.set(f"Drag across the graph to zoom the {mode_name}.")
        else:
            self.status_text.set("Area zoom cancelled.")

    def _on_area_zoom(self, start: MouseEvent, end: MouseEvent) -> None:
        if start.xdata is None or end.xdata is None or start.ydata is None or end.ydata is None:
            return
        x_limits = sorted((start.xdata, end.xdata))
        y_limits = sorted((start.ydata, end.ydata))
        zoom_x = self.zoom_axis_mode in {"both", "x"}
        zoom_y = self.zoom_axis_mode in {"both", "y"}
        if (zoom_x and x_limits[0] == x_limits[1]) or (zoom_y and y_limits[0] == y_limits[1]):
            return
        old_left = self.ax.get_ylim()
        old_right = self.ax_right.get_ylim()
        if zoom_x:
            self.ax.set_xlim((x_limits[0], x_limits[1]))
        if zoom_y:
            self.ax.set_ylim((y_limits[0], y_limits[1]))
        if zoom_y and not self.separate_graphs_mode and any(axis == "right" for _, axis in self._visible_signals()):
            left_span = old_left[1] - old_left[0]
            if left_span:
                low_fraction = (y_limits[0] - old_left[0]) / left_span
                high_fraction = (y_limits[1] - old_left[0]) / left_span
                right_span = old_right[1] - old_right[0]
                self.ax_right.set_ylim(
                    old_right[0] + low_fraction * right_span,
                    old_right[0] + high_fraction * right_span,
                )
        self.zoom_mode = False
        self.zoom_button.configure(style="Icon.TMenubutton")
        if self.area_selector is not None:
            self.area_selector.set_active(False)
        self.canvas.draw_idle()
        self.status_text.set("Zoomed to selected area.")

    def _on_scroll_zoom(self, event: MouseEvent) -> None:
        if not self._is_plot_axis(event.inaxes) or event.xdata is None or event.ydata is None:
            return
        step = getattr(event, "step", 0.0)
        if not step:
            button = getattr(event, "button", None)
            step = 1.0 if button == "up" else -1.0 if button == "down" else 0.0
        if not step:
            return
        modifiers = {str(value).casefold() for value in (getattr(event, "modifiers", ()) or ())}
        key = str(getattr(event, "key", "") or "").casefold()
        zoom_x = not ("control" in modifiers or "ctrl" in modifiers or key in {"control", "ctrl"})
        zoom_y = not ("shift" in modifiers)
        if key == "shift":
            zoom_y = False
        if key in {"control", "ctrl"}:
            zoom_x = False
        if not zoom_x and not zoom_y:
            return
        factor = math.pow(1.18, -max(-20.0, min(20.0, float(step))))
        if zoom_x:
            self.ax.set_xlim(self._zoom_limits(self.ax.get_xlim(), event.xdata, factor))
        if zoom_y:
            axis = event.inaxes
            if not isinstance(axis, Axes):
                return
            axis.set_ylim(self._zoom_limits(axis.get_ylim(), event.ydata, factor))
        self.canvas.draw_idle()

    @staticmethod
    def _zoom_limits(limits: tuple[float, float], anchor: float, factor: float) -> tuple[float, float]:
        lower, upper = limits
        return anchor - (anchor - lower) * factor, anchor + (upper - anchor) * factor

    def _toggle_live_cursor(self) -> None:
        if self.live_cursor_mode:
            self._deactivate_live_cursor()
            self.status_text.set("Live cursor disabled.")
            return
        self.live_cursor_mode = True
        self.live_cursor_seconds = None
        self._update_live_cursor_text()
        self.cursor_mode = None
        self.dragging_cursor = None
        self.zoom_mode = False
        self.zoom_button.configure(style="Icon.TMenubutton")
        self._update_cursor_menu_style()
        self.canvas_widget.configure(cursor="crosshair")
        if self.area_selector is not None:
            self.area_selector.set_active(False)
        self._refresh_measurements()
        self.status_text.set("Move the pointer over the graph to inspect signal values.")

    def _deactivate_live_cursor(self) -> None:
        self.live_cursor_mode = False
        self.live_cursor_seconds = None
        self._update_live_cursor_text()
        self._update_cursor_menu_style()
        self.canvas_widget.configure(cursor="arrow")
        for artist in self.live_cursor_artists:
            artist.set_visible(False)
        self._hide_live_popup()
        self._refresh_measurements()
        self.canvas.draw_idle()

    def _fit_all(self) -> None:
        if not self._visible_signals():
            self.status_text.set("Show at least one signal before fitting the graph.")
            return
        self._draw_plot(reset_view=True)
        self.status_text.set("Fit all axes to the visible signals.")

    def _fit_x(self) -> None:
        if not self._visible_signals():
            self.status_text.set("Show at least one signal before fitting the graph.")
            return
        y_limits = [axis.get_ylim() for axis in self.plot_axes]
        right_limits = [axis.get_ylim() for axis in self.right_plot_axes]
        self._draw_plot(reset_view=True)
        for axis, limits in zip(self.plot_axes, y_limits):
            axis.set_ylim(limits)
        for axis, limits in zip(self.right_plot_axes, right_limits):
            axis.set_ylim(limits)
        self.canvas.draw_idle()
        self.status_text.set("Fit X axis; Y ranges unchanged.")

    def _fit_y(self) -> None:
        if not self._visible_signals():
            self.status_text.set("Show at least one signal before fitting the graph.")
            return
        x_limits = self.ax.get_xlim()
        self._draw_plot(reset_view=True)
        self.ax.set_xlim(x_limits)
        self.canvas.draw_idle()
        self.status_text.set("Fit Y axes; X range unchanged.")

    def _display_samples(self, signal: Signal, point_limit: int) -> tuple[list[float], list[float]]:
        cached = self.display_cache.get(signal.uid)
        if cached is not None and cached[0] == point_limit:
            return cached[1], cached[2]
        sample_count = len(signal.times)
        if sample_count <= point_limit:
            result = (signal.times, signal.values)
        else:
            bucket_size = max(1, math.ceil(sample_count / max(1, point_limit // 2)))
            display_times: list[float] = []
            display_values: list[float] = []
            for bucket_start in range(0, sample_count, bucket_size):
                bucket_end = min(sample_count, bucket_start + bucket_size)
                index = bucket_start
                while index < bucket_end:
                    while index < bucket_end and not math.isfinite(signal.values[index]):
                        if not display_values or math.isfinite(display_values[-1]):
                            display_times.append(signal.times[index])
                            display_values.append(math.nan)
                        index += 1
                    segment_start = index
                    while index < bucket_end and math.isfinite(signal.values[index]):
                        index += 1
                    if segment_start == index:
                        continue
                    segment = range(segment_start, index)
                    minimum_index = min(segment, key=signal.values.__getitem__)
                    maximum_index = max(segment, key=signal.values.__getitem__)
                    for sample_index in sorted({segment_start, index - 1, minimum_index, maximum_index}):
                        display_times.append(signal.times[sample_index])
                        display_values.append(signal.values[sample_index])
            result = (display_times, display_values)
        self.display_cache[signal.uid] = (point_limit, result[0], result[1])
        return result

    def _on_canvas_resize(self, _event: object) -> None:
        if self.resize_job_id is None:
            self.resize_job_id = self.root.after(100, self._redraw_after_resize)

    def _redraw_after_resize(self) -> None:
        self.resize_job_id = None
        self._draw_plot(reset_view=False)

    def _draw_plot(self, reset_view: bool) -> None:
        visible_signals = self._visible_signals()
        previous_xlim = self.ax.get_xlim() if visible_signals and not reset_view else None
        if self.area_selector is not None:
            self.area_selector.set_active(False)
            self.area_selector.disconnect_events()
            self.area_selector = None
        self.figure.clear()
        self.plot_axes = []
        self.right_plot_axes = []
        self.signal_axes = {}
        self.cursor_artists = {}
        self.live_cursor_artists = []
        self.live_cursor_artist = None
        self.live_popup_artist = None
        self.live_popup_uid = None
        self.figure.patch.set_facecolor("#111a19")
        stacked = self.separate_graphs_mode and bool(visible_signals)
        if stacked:
            count = len(visible_signals)
            for index, (uid, _axis_name) in enumerate(visible_signals):
                axis = self.figure.add_subplot(
                    count,
                    1,
                    index + 1,
                    sharex=self.plot_axes[0] if self.plot_axes else None,
                )
                self.plot_axes.append(axis)
                self.signal_axes[uid] = axis
            self.ax = self.plot_axes[0]
            self.ax_right = self.ax
        else:
            self.ax = self.figure.add_subplot(111)
            self.ax_right = self.ax.twinx()
            self.plot_axes = [self.ax]
            self.right_plot_axes = [self.ax_right]
            for uid, axis_name in visible_signals:
                self.signal_axes[uid] = self.ax if axis_name == "left" else self.ax_right

        for axis in self.plot_axes + self.right_plot_axes:
            axis.set_facecolor("#111a19")
            axis.tick_params(colors="#a8bbb5", labelsize=8)
            for spine in axis.spines.values():
                spine.set_color("#42524e")
            axis.grid(True, color="#30403c", linewidth=0.7, alpha=0.72)
            axis.set_axisbelow(True)
        scale = 1000 if self.time_unit.get() == "ms" else 1
        if stacked:
            for index, (uid, _axis_name) in enumerate(visible_signals):
                axis = self.signal_axes[uid]
                axis.set_title(self.signals[uid].name, loc="left", color=self.signal_colors[uid], fontsize=8, pad=2)
                axis.set_ylabel("Value", color="#a8bbb5", fontsize=8)
                if index == len(visible_signals) - 1:
                    axis.set_xlabel(f"Time [{self.time_unit.get()}]", color="#c5d3cf", labelpad=6)
                else:
                    axis.tick_params(labelbottom=False)
        else:
            self.ax.set_xlabel(f"Time [{self.time_unit.get()}]", color="#c5d3cf", labelpad=9)
            self.ax.set_ylabel("Value", color="#48d6ad", labelpad=9)
            self.ax_right.set_ylabel("Value (right axis)", color="#ff9d72", labelpad=9)
            self.ax.yaxis.set_label_coords(-0.07, 0.5)
            self.ax_right.yaxis.set_label_position("right")
            self.ax_right.yaxis.tick_right()
            self.ax_right.yaxis.set_label_coords(1.07, 0.5)
            self.ax.yaxis.label.set_color("#48d6ad")
            self.ax_right.yaxis.label.set_color("#ff9d72")

        canvas_width = max(400, self.canvas_widget.winfo_width())
        point_limit = max(64, min(3000, canvas_width * 3 // max(1, len(visible_signals))))
        for uid, axis_name in visible_signals:
            signal = self.signals[uid]
            axis = self.signal_axes[uid]
            color = self.signal_colors[uid]
            distinct = len(set(signal.values[: min(512, len(signal.values))]))
            drawstyle = "steps-post" if distinct <= 16 and len(signal.values) > 20 else "default"
            display_times, display_values = self._display_samples(signal, point_limit)
            axis.plot(
                [value * scale for value in display_times],
                display_values,
                color=color,
                linewidth=1.15,
                drawstyle=drawstyle,
            )
            if self.samples_visible:
                axis.plot(
                    [value * scale for value in signal.times],
                    signal.values,
                    linestyle="None",
                    marker="o",
                    color=color,
                    markersize=2.4,
                    markerfacecolor=color,
                    markeredgewidth=0,
                    alpha=0.88,
                    zorder=4,
                )

        for cursor_name, seconds in self.cursor_seconds.items():
            self.cursor_artists[cursor_name] = [
                axis.axvline(
                    seconds * scale,
                    color="#f4eee1" if cursor_name == "A" else "#f3c75f",
                    linewidth=1.1,
                    linestyle="--",
                    zorder=5,
                )
                for axis in self.plot_axes
            ]
        self.live_cursor_artists = []
        if self.live_cursor_mode and self.live_cursor_seconds is not None:
            self.live_cursor_artists = [
                axis.axvline(
                    self.live_cursor_seconds * scale,
                    color="#85c9ff",
                    linewidth=1.2,
                    linestyle=":",
                    zorder=6,
                )
                for axis in self.plot_axes
            ]
        self.live_cursor_artist = self.live_cursor_artists[0] if self.live_cursor_artists else None
        if visible_signals:
            if previous_xlim:
                self.ax.set_xlim(previous_xlim)
        else:
            self.ax.text(0.5, 0.52, "OPEN A DUMP  /  SELECT SIGNALS TO SHOW", transform=self.ax.transAxes, ha="center", va="center", color="#80938d", fontsize=11, fontweight="bold")
            self.ax.set_xlim(0, 1)
            self.ax.set_ylim(0, 1)
        if stacked:
            self.figure.subplots_adjust(left=0.075, right=0.985, top=0.985, bottom=0.06, hspace=0.28)
        else:
            self.figure.subplots_adjust(left=0.085, right=0.915, top=0.985, bottom=0.075)
        self.area_selector = RectangleSelector(
            self.ax,
            self._on_area_zoom,
            useblit=True,
            button=[MouseButton.LEFT],
            minspanx=2,
            minspany=2,
            spancoords="pixels",
            props={"facecolor": "#48d6ad", "edgecolor": "#a5f3d6", "alpha": 0.18},
        )
        self.area_selector.set_active(self.zoom_mode)
        self.canvas.draw_idle()
        self._refresh_measurements()
        self._refresh_live_popup()

    def _arm_cursor(self, name: str) -> None:
        if self.live_cursor_mode:
            self._deactivate_live_cursor()
        self.cursor_mode = name
        self.dragging_cursor = None
        self.zoom_mode = False
        self.zoom_button.configure(style="Icon.TMenubutton")
        if self.area_selector is not None:
            self.area_selector.set_active(False)
        self.status_text.set(f"Cursor {name} armed. Click the graph to place it.")
        self._update_cursor_menu_style()

    def _on_plot_click(self, event: MouseEvent) -> None:
        x_position = event.xdata
        if not self._is_plot_axis(event.inaxes) or x_position is None:
            return
        scale = 1000 if self.time_unit.get() == "ms" else 1
        if self.cursor_mode is not None:
            self.cursor_seconds[self.cursor_mode] = x_position / scale
            placed = self.cursor_mode
            self.cursor_mode = None
            self._update_cursor_delta_text()
            self._update_cursor_menu_style()
            self._draw_plot(reset_view=False)
            self.status_text.set(f"Cursor {placed} placed. Drag its line to move it.")
            return
        if self.zoom_mode or self.live_cursor_mode or event.button != MouseButton.LEFT:
            return
        cursor_pixels = {
            name: self.ax.transData.transform((seconds * scale, 0))[0]
            for name, seconds in self.cursor_seconds.items()
        }
        if cursor_pixels:
            closest_name = min(cursor_pixels, key=lambda name: abs(cursor_pixels[name] - event.x))
            if abs(cursor_pixels[closest_name] - event.x) <= 10:
                self.dragging_cursor = closest_name
                self.canvas_widget.configure(cursor="fleur")
                self.status_text.set(f"Moving cursor {closest_name}.")

    def _on_plot_motion(self, event: MouseEvent) -> None:
        x_position = event.xdata
        if self.dragging_cursor is not None:
            if self._is_plot_axis(event.inaxes) and x_position is not None:
                scale = 1000 if self.time_unit.get() == "ms" else 1
                seconds = x_position / scale
                self.cursor_seconds[self.dragging_cursor] = seconds
                self._update_cursor_delta_text()
                for artist in self.cursor_artists.get(self.dragging_cursor, []):
                    artist.set_xdata((x_position, x_position))
                self._schedule_measurement_refresh()
                self.canvas.draw_idle()
            return
        if not self.live_cursor_mode:
            return
        if not self._is_plot_axis(event.inaxes) or x_position is None:
            if self.live_cursor_seconds is not None:
                self.live_cursor_seconds = None
                self._update_live_cursor_text()
                for artist in self.live_cursor_artists:
                    artist.set_visible(False)
                self._hide_live_popup()
                self._schedule_measurement_refresh()
                self.canvas.draw_idle()
            return

        scale = 1000 if self.time_unit.get() == "ms" else 1
        self.live_cursor_seconds = x_position / scale
        self._update_live_cursor_text()
        if not self.live_cursor_artists:
            self.live_cursor_artists = [
                axis.axvline(
                    x_position,
                    color="#85c9ff",
                    linewidth=1.2,
                    linestyle=":",
                    zorder=6,
                )
                for axis in self.plot_axes
            ]
        else:
            for artist in self.live_cursor_artists:
                artist.set_xdata((x_position, x_position))
                artist.set_visible(True)
        self.live_cursor_artist = self.live_cursor_artists[0] if self.live_cursor_artists else None
        self._refresh_live_popup()
        self._schedule_measurement_refresh()
        self.canvas.draw_idle()

    def _on_plot_release(self, _event: MouseEvent) -> None:
        if self.dragging_cursor is None:
            return
        cursor_name = self.dragging_cursor
        self.dragging_cursor = None
        self.canvas_widget.configure(cursor="crosshair" if self.live_cursor_mode else "arrow")
        scale = 1000 if self.time_unit.get() == "ms" else 1
        position = self.cursor_seconds[cursor_name] * scale
        self.status_text.set(f"Cursor {cursor_name} moved to {position:.6g} {self.time_unit.get()}.")

    def _schedule_measurement_refresh(self) -> None:
        if self.measurement_refresh_id is None:
            self.measurement_refresh_id = self.root.after(40, self._flush_measurement_refresh)

    def _flush_measurement_refresh(self) -> None:
        self.measurement_refresh_id = None
        self._refresh_measurements()

    def _on_time_unit_changed(self, *_args: str) -> None:
        self._draw_plot(reset_view=False)
        self._update_live_cursor_text()
        self._update_cursor_delta_text()

    def _update_live_cursor_text(self) -> None:
        if not self.live_cursor_mode or self.live_cursor_seconds is None:
            self.live_cursor_text.set("t: --")
            return
        scale = 1000 if self.time_unit.get() == "ms" else 1
        position = self.live_cursor_seconds * scale
        self.live_cursor_text.set(f"t: {position:.6g} {self.time_unit.get()}")

    def _update_cursor_delta_text(self) -> None:
        if "A" not in self.cursor_seconds or "B" not in self.cursor_seconds:
            self.cursor_delta_text.set("Δt: --")
            return
        scale = 1000 if self.time_unit.get() == "ms" else 1
        delta = (self.cursor_seconds["B"] - self.cursor_seconds["A"]) * scale
        self.cursor_delta_text.set(f"Δt: {delta:.6g} {self.time_unit.get()}")

    def _clear_cursors(self) -> None:
        self.cursor_seconds.clear()
        self._update_cursor_delta_text()
        self.cursor_mode = None
        self.dragging_cursor = None
        self._update_cursor_menu_style()
        self._draw_plot(reset_view=False)
        self.status_text.set("Cursors cleared.")

    def _sample_at(self, signal: Signal, seconds: float) -> float | None:
        value = signal.values[self._nearest_sample_index(signal, seconds)]
        return value if math.isfinite(value) else None

    @staticmethod
    def _nearest_sample_index(signal: Signal, seconds: float) -> int:
        index = bisect.bisect_left(signal.times, seconds, 0, len(signal.times), key=None)
        if index == 0:
            candidates = (0,)
        elif index >= len(signal.times):
            candidates = (len(signal.times) - 1,)
        else:
            candidates = (index - 1, index)
        return min(candidates, key=lambda item: abs(signal.times[item] - seconds))

    def _refresh_measurements(self) -> None:
        tree = self.signal_tree
        if tree is None:
            return
        for uid, axis in self.plotted:
            signal = self.signals[uid]
            visible = self.signal_visibility[uid].get()
            value_live = (
                self._sample_at(signal, self.live_cursor_seconds)
                if visible and self.live_cursor_mode and self.live_cursor_seconds is not None
                else None
            )
            value_a = self._sample_at(signal, self.cursor_seconds["A"]) if visible and "A" in self.cursor_seconds else None
            value_b = self._sample_at(signal, self.cursor_seconds["B"]) if visible and "B" in self.cursor_seconds else None
            delta = value_b - value_a if value_a is not None and value_b is not None else None
            item_values = (
                "L" if axis == "left" else "R",
                "✓" if visible else "",
                *(self._format_number(value) for value in (value_live, value_a, value_b, delta)),
            )
            if tree.exists(uid):
                tree.item(uid, values=item_values)

    @staticmethod
    def _format_number(value: float | None) -> str:
        return "" if value is None else f"{value:.7g}"

    def _save_graph(self) -> None:
        path = filedialog.asksaveasfilename(
            title="Save graph image",
            defaultextension=".png",
            filetypes=(("PNG image", "*.png"), ("PDF document", "*.pdf"), ("SVG image", "*.svg")),
        )
        if not path:
            return
        legend = None
        export_title = None
        subplot_parameters = {
            name: getattr(self.figure.subplotpars, name)
            for name in ("left", "right", "bottom", "top", "wspace", "hspace")
        }
        try:
            file_title = " · ".join(Path(dataset[0]).name for dataset in self.datasets)
            export_title = self.figure.suptitle(file_title or "DaqIdea Graph Panel", y=0.995, color=TEXT, fontsize=10)
            visible_signals = self._visible_signals()
            if visible_signals and not self.separate_graphs_mode:
                handles = [
                    Line2D([], [], color=self.signal_colors[uid], linewidth=1.8)
                    for uid, _axis in visible_signals
                ]
                labels = [self.signals[uid].label for uid, _axis in visible_signals]
                legend_columns = min(4, len(handles))
                legend_rows = math.ceil(len(handles) / legend_columns)
                self.figure.subplots_adjust(top=max(0.55, 0.92 - 0.025 * legend_rows))
                legend = self.figure.legend(
                    handles,
                    labels,
                    loc="upper center",
                    bbox_to_anchor=(0.5, 0.955),
                    ncol=legend_columns,
                    fontsize=8,
                    facecolor=PANEL_BG,
                    edgecolor="#42524e",
                    labelcolor=TEXT,
                    framealpha=0.98,
                )
            self.figure.savefig(path, dpi=180, facecolor=self.figure.get_facecolor(), bbox_inches="tight")
            self.status_text.set(f"Graph saved: {Path(path).name}")
        except Exception as error:
            messagebox.showerror("Could not save graph", str(error), parent=self.root)
        finally:
            if legend is not None:
                legend.remove()
            if export_title is not None:
                export_title.remove()
            self.figure.subplots_adjust(**subplot_parameters)
            self.canvas.draw_idle()

    def _export_csv(self) -> None:
        visible_signals = self._visible_signals()
        if not visible_signals:
            messagebox.showinfo("Export CSV", "Add at least one signal to the graph first.", parent=self.root)
            return
        path = filedialog.asksaveasfilename(
            title="Export plotted signals",
            defaultextension=".csv",
            filetypes=(("CSV file", "*.csv"),),
        )
        if not path:
            return
        try:
            signals = [self.signals[uid] for uid, _axis in visible_signals]
            headers = []
            for signal in signals:
                headers.extend((f"{signal.label} time (s)", f"{signal.label} value"))
            with open(path, "w", newline="", encoding="utf-8-sig") as output:
                writer = csv.writer(output)
                writer.writerow(headers)
                for row_index in range(max(len(signal.times) for signal in signals)):
                    row = []
                    for signal in signals:
                        if row_index < len(signal.times):
                            value = signal.values[row_index]
                            row.extend((signal.times[row_index], "" if not math.isfinite(value) else value))
                        else:
                            row.extend(("", ""))
                    writer.writerow(row)
            self.status_text.set(f"CSV exported: {Path(path).name}")
        except Exception as error:
            messagebox.showerror("Could not export CSV", str(error), parent=self.root)


def main() -> None:
    root = TkinterDnD.Tk()
    DaqIdeaPlotter(root)
    root.mainloop()


if __name__ == "__main__":
    main()